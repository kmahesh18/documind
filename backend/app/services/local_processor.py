"""
Local file processor - Supports ALL document types.
- PDF, DOCX, XLSX, TXT: Direct text extraction
- Images: Gemini Vision for OCR/description
- Audio/Video: Groq Whisper for transcription
"""
import httpx
import io
import tempfile
import os
from app.core.config import get_settings

settings = get_settings()


class LocalFileProcessor:
    """
    Process files locally.
    - Documents: pdfplumber, python-docx, openpyxl
    - Images: Gemini Vision API
    - Audio/Video: Groq Whisper API
    """
    
    def __init__(self):
        self._rag_service = None
    
    def _get_rag_service(self):
        """Lazy load RAG service."""
        if self._rag_service is None:
            from app.services.langchain_rag import get_rag_service
            self._rag_service = get_rag_service()
        return self._rag_service
    
    async def process_file(
        self, 
        file_url: str, 
        file_id: str, 
        user_id: str, 
        filename: str,
        file_type: str
    ) -> bool:
        """Process any file type and store in vector DB."""
        try:
            print(f"Processing: {filename} ({file_type})")
            
            # Download file
            content = await self._download_file(file_url)
            if not content:
                print(f"Failed to download: {filename}")
                return False
            
            # Extract text based on type
            text = await self._extract_text(content, file_type, filename)
            if not text or text.strip() == "":
                print(f"No text extracted from: {filename}")
                return False
            
            print(f"Extracted {len(text)} chars from {filename}")
            
            # Store in vector DB
            rag = self._get_rag_service()
            num_vectors = await rag.process_and_store_document(
                text=text,
                file_id=file_id,
                user_id=user_id,
                filename=filename,
                file_type=file_type,
                file_url=file_url
            )
            
            print(f"Stored {num_vectors} vectors for {filename}")
            return True
            
        except Exception as e:
            print(f"Error processing {filename}: {e}")
            import traceback
            traceback.print_exc()
            return False
    
    async def _download_file(self, url: str) -> bytes:
        """Download file from URL."""
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(url, timeout=120.0)
                response.raise_for_status()
                return response.content
        except Exception as e:
            print(f"Download error: {e}")
            return None
    
    async def _extract_text(self, content: bytes, file_type: str, filename: str) -> str:
        """Extract text from any file type."""
        try:
            print(f"[Extract] Processing {filename} as type: {file_type}")
            
            if file_type in ["txt", "csv"]:
                text = content.decode("utf-8", errors="ignore")
                print(f"[Extract] Decoded text file: {len(text)} chars")
                return text
            
            elif file_type == "pdf":
                return self._extract_pdf(content)
            
            elif file_type == "docx":
                return self._extract_docx(content)
            
            elif file_type == "xlsx":
                return self._extract_xlsx(content)
            
            elif file_type in ["ppt", "pptx"]:
                return self._extract_pptx(content)
            
            elif file_type == "image":
                return await self._extract_image(content, filename)
            
            elif file_type in ["audio", "video"]:
                return await self._transcribe_media(content, filename, file_type)
            
            else:
                # Try as text
                text = content.decode("utf-8", errors="ignore")
                print(f"[Extract] Fallback text decode: {len(text)} chars")
                return text
                
        except Exception as e:
            print(f"[Extract] Error for {file_type}: {e}")
            import traceback
            traceback.print_exc()
            return ""
    
    def _extract_pdf(self, content: bytes) -> str:
        """Extract text from PDF."""
        try:
            import pdfplumber
            
            text_parts = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(f"[Page {i}]\n{page_text}")
            
            return "\n\n".join(text_parts)
        except Exception as e:
            print(f"PDF error: {e}")
            return ""
    
    def _extract_docx(self, content: bytes) -> str:
        """Extract text from DOCX."""
        try:
            from docx import Document
            
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except Exception as e:
            print(f"DOCX error: {e}")
            return ""
    
    def _extract_xlsx(self, content: bytes) -> str:
        """Extract text from Excel."""
        try:
            import openpyxl
            
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            text_parts = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                text_parts.append(f"[Sheet: {sheet_name}]")
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) if cell else "" for cell in row)
                    if row_text.strip(" |"):
                        text_parts.append(row_text)
            
            return "\n".join(text_parts)
        except Exception as e:
            print(f"Excel error: {e}")
            return ""

    def _extract_pptx(self, content: bytes) -> str:
        """Extract text from PPTX slides using python-pptx."""
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(content))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                slide_text_parts = []
                for shape in slide.shapes:
                    # Many shapes may not have text_frame
                    if hasattr(shape, "text"):
                        text = getattr(shape, "text")
                        if text and text.strip():
                            slide_text_parts.append(text.strip())

                # Slide notes
                try:
                    notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide is not None else ""
                    if notes and notes.strip():
                        slide_text_parts.append(f"[Notes] {notes.strip()}")
                except Exception:
                    pass

                if slide_text_parts:
                    parts.append(f"[Slide {i}]\n" + "\n".join(slide_text_parts))

            return "\n\n".join(parts)
        except Exception as e:
            print(f"PPTX extraction error: {e}")
            return ""
    
    async def _extract_image(self, content: bytes, filename: str) -> str:
        """Extract text from image using Gemini Vision."""
        try:
            import google.generativeai as genai
            import base64
            
            print(f"[Image] Starting OCR for: {filename}, size: {len(content)} bytes")
            
            genai.configure(api_key=settings.gemini_api_key)
            model = genai.GenerativeModel("gemini-2.0-flash")
            
            # Determine mime type
            ext = filename.lower().split(".")[-1]
            mime_map = {
                "png": "image/png",
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
                "gif": "image/gif",
                "webp": "image/webp",
                "bmp": "image/bmp"
            }
            mime_type = mime_map.get(ext, "image/png")
            print(f"[Image] MIME type: {mime_type}")
            
            # Create image part with base64 encoding
            image_part = {
                "mime_type": mime_type,
                "data": base64.b64encode(content).decode("utf-8")
            }
            
            # Ask Gemini to extract all text and describe
            prompt = """Analyze this image thoroughly:

1. Extract ALL visible text exactly as shown (OCR) - be thorough and precise
2. Describe what the image contains in detail
3. Note any important visual elements, diagrams, charts, or data

Be comprehensive - extract every piece of text and visual information you can see."""

            response = model.generate_content([prompt, image_part])
            
            if response and response.text:
                result = f"[Image: {filename}]\n\n{response.text}"
                print(f"[Image] OCR done: {len(result)} chars extracted")
                return result
            else:
                print(f"[Image] No response from Gemini for {filename}")
                return f"[Image file: {filename}] - No text could be extracted"
            
        except Exception as e:
            print(f"[Image] Extraction error for {filename}: {e}")
            import traceback
            traceback.print_exc()
            return f"[Image file: {filename}] - OCR failed: {str(e)}"
    
    async def _transcribe_media(self, content: bytes, filename: str, file_type: str) -> str:
        """Transcribe audio/video using Groq Whisper API."""
        try:
            if not settings.groq_api_key:
                print("Groq API key not set!")
                return f"[{file_type.title()} file: {filename}] - No transcription (API key missing)"
            
            ext = filename.lower().split(".")[-1]
            
            # Write to temp file
            with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            
            try:
                print(f"Transcribing {filename} with Groq Whisper...")
                
                async with httpx.AsyncClient() as client:
                    with open(tmp_path, "rb") as f:
                        files = {"file": (filename, f)}
                        data = {
                            "model": "whisper-large-v3",
                            "response_format": "verbose_json"
                        }
                        
                        response = await client.post(
                            "https://api.groq.com/openai/v1/audio/transcriptions",
                            headers={"Authorization": f"Bearer {settings.groq_api_key}"},
                            files=files,
                            data=data,
                            timeout=300.0
                        )
                        response.raise_for_status()
                        result = response.json()
                
                transcript = result.get("text", "")
                
                if transcript:
                    segments = result.get("segments", [])
                    if segments:
                        # Format with timestamps
                        parts = [f"[{file_type.title()}: {filename}]\n\n[Transcript]"]
                        for seg in segments:
                            start = seg.get("start", 0)
                            text = seg.get("text", "").strip()
                            if text:
                                mins, secs = divmod(int(start), 60)
                                parts.append(f"[{mins:02d}:{secs:02d}] {text}")
                        result_text = "\n".join(parts)
                    else:
                        result_text = f"[{file_type.title()}: {filename}]\n\n[Transcript]\n{transcript}"
                    
                    print(f"Transcription done: {len(result_text)} chars")
                    return result_text
                else:
                    return f"[{file_type.title()}: {filename}] (no speech detected)"
                    
            finally:
                os.unlink(tmp_path)
                
        except Exception as e:
            print(f"Transcription error: {e}")
            import traceback
            traceback.print_exc()
            return f"[{file_type.title()} file: {filename}] - Transcription failed"


# Singleton
local_processor = LocalFileProcessor()
